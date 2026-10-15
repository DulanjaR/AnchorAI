"""Apply review verdicts and split into train / val / test BY TASK.

All five tiers of a task land in the same split, so the test set measures tier behaviour on
tasks the model has never seen (splitting by example would leak the task text).
Splits are stratified by module.

    python -m components.c4_hcde.training.split_dataset

Rules:
- verdict ``reject`` -> dropped; ``fix`` -> reply replaced by ``fixed_reply``; ``keep`` -> kept.
- flagged examples (compliance failure or disclosure outside its band) with no verdict are
  dropped (pass --keep-unreviewed-flagged to keep them). A fixed reply loses its disclosure
  score; rerun label_disclosure if you need it.
"""
from __future__ import annotations

import argparse
import csv
import json
import random
from collections import Counter, defaultdict
from pathlib import Path

from .build_dataset import DEFAULT_OUT, load_jsonl
from .compliance import check
from .review_sample import DEFAULT_REVIEW, needs_review


def load_verdicts(path: Path) -> dict[tuple[str, str], dict]:
    if not path.exists():
        return {}
    with path.open(newline="") as f:
        return {(r["task_id"], r["tier"]): r for r in csv.DictReader(f) if r.get("verdict", "").strip()}


def load_decisions(path: Path) -> dict[tuple[str, str], dict]:
    """Extra keep/reject decisions made outside the review sheet (one JSON object per line:
    task_id, tier, verdict, reason, by). The review sheet wins when both cover the same reply."""
    if not path.exists():
        return {}
    with path.open() as f:
        rows = [json.loads(line) for line in f if line.strip()]
    out = {(r["task_id"], r["tier"]): r for r in rows}
    # The researcher's own verdicts in review_verdicts.xlsx (column "YOUR verdict") override Claude's.
    xlsx = path.parent / "review_verdicts.xlsx"
    if xlsx.exists():
        try:
            from openpyxl import load_workbook
            ws = load_workbook(xlsx, read_only=True)["Verdicts"]
            head = [str(c.value or "") for c in next(ws.iter_rows(max_row=1))]
            i_v = next(i for i, h in enumerate(head) if h.startswith("YOUR verdict"))
            i_n = i_v + 1
            for row in ws.iter_rows(min_row=2, values_only=True):
                v = str(row[i_v] or "").strip().lower()
                if v in ("keep", "reject"):
                    out[(row[0], row[1])] = {"task_id": row[0], "tier": row[1], "verdict": v,
                                             "reason": row[i_n] or "", "by": "researcher"}
        except ImportError:
            print("openpyxl not installed: ignoring review_verdicts.xlsx")
    return out


def split_tasks(task_modules: dict[str, str], ratios: tuple[float, float, float], seed: int) -> dict[str, str]:
    by_module: dict[str, list[str]] = defaultdict(list)
    for tid, mod in sorted(task_modules.items()):
        by_module[mod].append(tid)
    rng = random.Random(seed)
    assignment: dict[str, str] = {}
    for mod in sorted(by_module):
        ids = by_module[mod]
        rng.shuffle(ids)
        n_test = max(1, round(len(ids) * ratios[2]))
        n_val = max(1, round(len(ids) * ratios[1]))
        for i, tid in enumerate(ids):
            assignment[tid] = "test" if i < n_test else "val" if i < n_test + n_val else "train"
    return assignment


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--raw", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--review", type=Path, default=DEFAULT_REVIEW)
    ap.add_argument("--out-dir", type=Path, default=DEFAULT_OUT.parent)
    ap.add_argument("--decisions", type=Path, default=DEFAULT_OUT.parent / "decisions.jsonl",
                    help="extra keep/reject decisions with reasons (applied when the review sheet has none)")
    ap.add_argument("--ratios", type=float, nargs=3, default=(0.75, 0.10, 0.15))
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--keep-unreviewed-flagged", action="store_true")
    args = ap.parse_args(argv)

    rows = load_jsonl(args.raw)
    verdicts = load_verdicts(args.review)
    decisions = load_decisions(args.decisions)
    kept, dropped = [], Counter()
    norm = lambda t: " ".join((t or "").split())
    stale = 0
    for r in rows:
        v = verdicts.get((r["task_id"], r["tier"]))
        # A verdict only counts for the exact reply that was reviewed; if the reply was regenerated
        # since, the old verdict is ignored.
        if v and v.get("reply") and norm(v["reply"]) != norm(r["messages"][2]["content"]):
            v, stale = None, stale + 1
        if not v and (r["task_id"], r["tier"]) in decisions:
            v = decisions[(r["task_id"], r["tier"])]
        verdict = (v or {}).get("verdict", "").strip().lower()
        if verdict == "reject":
            dropped["rejected in review"] += 1
            continue
        if verdict == "fix":
            fixed = (v.get("fixed_reply") or "").strip()
            if not fixed:
                dropped["fix with empty fixed_reply"] += 1
                continue
            r["messages"][2]["content"] = fixed
            r["compliance"] = check(r["tier"], fixed, r["messages"][1]["content"]).to_dict()
            r.pop("disclosure", None)
            r["reviewed"] = "fixed"
        elif verdict == "keep":
            r["reviewed"] = "kept"
        elif needs_review(r) and not args.keep_unreviewed_flagged:
            dropped["flagged, not reviewed"] += 1
            continue
        kept.append(r)

    assignment = split_tasks({r["task_id"]: r["module"] for r in kept}, tuple(args.ratios), args.seed)
    buckets: dict[str, list[dict]] = defaultdict(list)
    for r in kept:
        buckets[assignment[r["task_id"]]].append(r)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    for name in ("train", "val", "test"):
        with (args.out_dir / f"{name}.jsonl").open("w") as f:
            for r in sorted(buckets[name], key=lambda x: (x["task_id"], x["tier"])):
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        tiers = Counter(r["tier"] for r in buckets[name])
        n_tasks = len({r["task_id"] for r in buckets[name]})
        print(f"{name:5s}: {len(buckets[name]):4d} examples from {n_tasks:3d} tasks  {dict(sorted(tiers.items()))}")
    print(f"kept {len(kept)} of {len(rows)}; dropped {dict(dropped) or 0}; "
          f"{stale} review verdicts ignored because the reply was regenerated")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
