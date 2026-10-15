import io
import json
from pathlib import Path

import pytest

from components.c3_checkpointing.app.difficulty import (
    DEPTH, duration_points, format_minutes, predict_difficulty,
)
from components.c3_checkpointing.app.summary import build_summary, summary_to_text
from components.c3_checkpointing.parser.text_files import (
    analyze_docx, analyze_file, analyze_text, suggest_from_file,
)

VECTORS = json.loads((Path(__file__).parent / "vectors.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("case", VECTORS["predict"])
def test_difficulty_matches_shared_vectors(case):
    i = case["in"]
    r = predict_difficulty(i["minutes"], i["complexity"], i["openItems"], i["type"])
    assert r["score"] == case["score"]
    assert r["level"] == case["level"]
    assert sum(x["points"] for x in r["reasons"]) == r["score"]


@pytest.mark.parametrize("minutes,expected", VECTORS["format_minutes"])
def test_format_minutes_matches_shared_vectors(minutes, expected):
    assert format_minutes(minutes) == expected


def test_duration_boundaries():
    assert [duration_points(m) for m in (14, 15, 59, 60, 239, 240)] == [0, 1, 1, 2, 2, 3]
    assert duration_points(None) == 0 and duration_points(-3) == 0 and duration_points("abc") == 0


def test_monotonic_in_duration_and_complexity():
    scores = [predict_difficulty(m, "moderate", 2, "meeting")["score"] for m in (1, 10, 30, 90, 300, 2000)]
    assert scores == sorted(scores)
    scores = [predict_difficulty(60, c, 2, "meeting")["score"] for c in ("simple", "moderate", "complex")]
    assert scores == sorted(scores)


def test_depth_grows_with_level():
    assert [len(DEPTH[l]["suggested"]) for l in ("Low", "Medium", "High")] == [2, 4, 6]


@pytest.mark.parametrize("level", ["Low", "Medium", "High"])
def test_summary_matches_shared_vectors(level):
    v = VECTORS["summary"]
    s = build_summary(v["task"], v["checkpoint"], level, 300)
    assert [l["label"] for l in s["lines"]] == v["labels"][level]
    assert s["intro"] == v["intro_300"]


def test_summary_only_restates_student_text():
    v = VECTORS["summary"]
    s = build_summary(v["task"], v["checkpoint"], "High", 300)
    cp = v["checkpoint"]
    allowed = [x for x in cp.values() if isinstance(x, str)] + ["; ".join(cp["openItems"])]
    for line in s["lines"]:
        assert line["text"] in allowed


def test_summary_skips_empty_fields_and_handles_nothing_written():
    s = build_summary("T", {"workingPoint": "here", "nextStep": "  "}, "Medium", 10)
    assert [l["label"] for l in s["lines"]] == ["You stopped at"]
    assert "(No checkpoint text was written.)" in summary_to_text(build_summary("T", {}, "Low", 5))


HTML = "\n".join([
    "<html><head><title>Contact Us</title></head><body>",
    "<h1>Contact <b>Us</b></h1>",
    "<!-- TODO: validate email -->",
    "<script>function validateName(v){return v.length>1}",
    "// FIXME: show error",
    "const send = (a) => a;</script></body></html>",
])


def test_html_analysis_matches_the_javascript_version():
    r = analyze_text("contact.html", HTML)
    assert r["supported"]
    assert r["headings"] == ["Contact Us"]
    assert [t["text"] for t in r["todos"]] == ["validate email", "show error"]
    assert r["symbols"] == ["validateName", "send"]
    assert r["openCount"] == 2


def test_markdown_python_and_unsupported():
    md = analyze_text("notes.md", "# Report\n- [x] ER diagram\n- [ ] SQL queries\n- [ ] Conclusion\nTODO: references")
    assert (len(md["doneItems"]), len(md["openItems"]), md["openCount"]) == (1, 2, 3)
    py = analyze_text("clean.py", "def load(path):\n    pass\nclass Cleaner:\n    # TODO handle nulls\n    pass\n")
    assert py["symbols"] == ["load", "Cleaner"] and len(py["todos"]) == 1
    assert analyze_file("photo.png", b"x")["supported"] is False
    assert analyze_file("slides.pptx", b"x")["supported"] is False


def test_analysis_never_returns_the_file_text():
    secret = "SECRET_PASSWORD_123"
    r = analyze_file("a.txt", f"line one\n{secret}\n".encode())
    assert secret not in json.dumps(r)


def test_suggestions_come_from_the_file():
    s = suggest_from_file(analyze_text("contact.html", HTML))
    assert s["resources"] == "contact.html"
    assert "validate email" in s["workingPoint"]
    assert "validateName" in s["completed"]


def _docx_bytes():
    from docx import Document

    d = Document()
    d.add_heading("Database assignment report", level=1)
    d.add_paragraph("[x] ER diagram")
    d.add_paragraph("[ ] SQL queries for the reports section")
    d.add_paragraph("[ ] Conclusion")
    d.add_paragraph("TODO: add references in IEEE format")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_support():
    r = analyze_file("report.docx", _docx_bytes())
    assert r["supported"] and r["ext"] == "docx"
    assert r["headings"] == ["Database assignment report"]
    assert r["doneItems"] == ["ER diagram"]
    assert r["openItems"] == ["SQL queries for the reports section", "Conclusion"]
    assert r["openCount"] == 3


def test_corrupt_docx_is_reported_not_crashed():
    r = analyze_docx("bad.docx", b"this is not a zip file")
    assert r["supported"] is False


def test_oversized_file_is_refused():
    assert analyze_file("big.txt", b"a" * (2 * 1024 * 1024 + 1))["supported"] is False
