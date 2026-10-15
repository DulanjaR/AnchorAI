import csv
import io
import json
import random

import pytest

fastapi = pytest.importorskip("fastapi")
pytest.importorskip("httpx")
pytest.importorskip("multipart")
pytest.importorskip("docx")

from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import create_engine, select  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

from components.c3_checkpointing.app import models as m  # noqa: E402
from components.c3_checkpointing.app.generators import (  # noqa: E402
    LoraGenerator, build_model_input, check_model_output, generate_summary, to_training_fields,
)
from components.c3_checkpointing.app.main import app  # noqa: E402
from components.c3_checkpointing.model2.generate_data import generate_split  # noqa: E402
from shared.config import settings  # noqa: E402
from shared.db import Base, get_db  # noqa: E402

KEY = {"X-API-Key": settings.api_key}


@pytest.fixture()
def env():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def override():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override
    yield TestClient(app), Session
    app.dependency_overrides.clear()


def H(user="u1"):
    return {**KEY, "X-User-Id": user}


def tasks(client, user="u1"):
    return {t["title"]: t for t in client.get("/api/tasks", headers=H(user)).json()}


CP = {"workingPoint": "contact.html form has no submit handler", "nextStep": "Add submit button",
      "goal": "Finish the booking site", "problem": "How to validate the email?", "completed": "login.html",
      "resources": "Figma link", "openItems": ["validate email", "add submit"], "files": ["contact.html"]}


def interrupt(client, title="Hotel booking website", user="u1", **over):
    t = tasks(client, user)[title]
    body = {**CP, "type": "meeting", "minutes": 30, "complexity": "moderate", **over}
    return t, client.post(f"/api/tasks/{t['id']}/checkpoint", json=body, headers=H(user))


# ----------------------------------------------------------------- access
def test_health_is_public_and_lists_template_mode(env):
    client, _ = env
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json()["modes"][0] == "template"


def test_api_key_is_required(env):
    client, _ = env
    assert client.get("/api/tasks").status_code == 401
    assert client.get("/api/tasks", headers={"X-API-Key": "wrong"}).status_code == 401


def test_user_id_must_be_a_safe_pseudonym(env):
    client, _ = env
    assert client.get("/api/tasks", headers={**KEY, "X-User-Id": "ann smith@mail.com"}).status_code == 400


def test_each_user_has_their_own_tasks(env):
    client, _ = env
    a, b = tasks(client, "alice"), tasks(client, "bob")
    assert len(a) == len(b) == 4 and a["Hotel booking website"]["id"] != b["Hotel booking website"]["id"]
    assert client.patch(f"/api/tasks/{a['Hotel booking website']['id']}", json={"column": "done"}, headers=H("bob")).status_code == 404


# ----------------------------------------------------------------- predict
def test_predict_matches_the_shared_rule(env):
    client, _ = env
    r = client.post("/api/predict", json={"minutes": 300, "complexity": "complex", "openItems": 6, "type": "other_task"}, headers=KEY)
    body = r.json()
    assert (body["score"], body["level"], body["depth"]) == (8, "High", "Full")
    assert sum(x["points"] for x in body["reasons"]) == 8 and len(body["suggested"]) == 6


def test_predict_rejects_bad_input(env):
    client, _ = env
    assert client.post("/api/predict", json={"minutes": -5}, headers=KEY).status_code == 422
    assert client.post("/api/predict", json={"complexity": "huge"}, headers=KEY).status_code == 422


# ----------------------------------------------------------------- checkpoint flow
def test_full_interrupt_and_resume_flow(env):
    client, Session = env
    t, r = interrupt(client, minutes=300, type="other_task", complexity="complex", openItemsCount=6)
    assert r.status_code == 201
    out = r.json()
    assert out["column"] == "interrupted" and out["checkpoint"]["level"] == "High" and out["checkpoint"]["depth"] == "Full"

    s = client.get(f"/api/tasks/{t['id']}/summary", params={"awayMinutes": 300}, headers=H()).json()
    assert s["mode"] == "template" and s["level"] == "High" and len(s["lines"]) == 7 and "5 h" in s["intro"]

    r = client.post(f"/api/tasks/{t['id']}/resume", json={"action": "accepted", "viewedSeconds": 8, "awayMinutes": 300,
                                                          "shownText": s["text"]}, headers=H())
    assert r.status_code == 200 and r.json()["column"] == "doing" and r.json()["checkpoint"] is None

    assert client.get(f"/api/tasks/{t['id']}/summary", headers=H()).status_code == 409  # nothing left to resume
    events = [e["event"] for e in client.get("/api/events", headers=H()).json()]
    assert events[-2:] == ["checkpoint_saved", "summary_accepted"]
    with Session() as db:
        res = db.scalars(select(m.Resumption)).all()
        assert len(res) == 1 and res[0].action == "accepted" and res[0].viewed_seconds == 8 and res[0].summary_text == s["text"]


def test_depth_changes_the_summary_length(env):
    client, _ = env
    t, _ = interrupt(client, minutes=5, type="quick_message", complexity="simple", openItemsCount=0)
    low = client.get(f"/api/tasks/{t['id']}/summary", headers=H()).json()
    assert low["level"] == "Low" and [l["label"] for l in low["lines"]] == ["You stopped at", "Next step"]


def test_checkpoint_rules(env):
    client, _ = env
    todo = tasks(client)["Python data-cleaning script"]
    r = client.post(f"/api/tasks/{todo['id']}/checkpoint", json={**CP}, headers=H())
    assert r.status_code == 409  # only tasks in progress can be interrupted
    t = tasks(client)["Hotel booking website"]
    assert client.post(f"/api/tasks/{t['id']}/checkpoint", json={"type": "meeting"}, headers=H()).status_code == 422
    assert client.post("/api/tasks/nope/checkpoint", json={**CP}, headers=H()).status_code == 404
    too_long = {**CP, "goal": "x" * 5000}
    assert client.post(f"/api/tasks/{t['id']}/checkpoint", json=too_long, headers=H()).status_code == 422


def test_move_rules(env):
    client, _ = env
    t, _ = interrupt(client)
    assert client.patch(f"/api/tasks/{tasks(client)['Database assignment report']['id']}", json={"column": "interrupted"}, headers=H()).status_code == 409
    assert client.patch(f"/api/tasks/{t['id']}", json={"column": "doing"}, headers=H()).status_code == 409
    assert client.patch(f"/api/tasks/{t['id']}", json={"column": "todo"}, headers=H()).status_code == 409
    assert client.patch(f"/api/tasks/{t['id']}", json={"column": "done"}, headers=H()).json()["column"] == "done"


def test_edited_summary_is_logged_with_its_length(env):
    client, _ = env
    t, _ = interrupt(client)
    assert client.post(f"/api/tasks/{t['id']}/resume", json={"action": "edited", "finalText": " "}, headers=H()).status_code == 422
    r = client.post(f"/api/tasks/{t['id']}/resume", json={"action": "edited", "finalText": "My own version of the summary",
                                                          "shownText": "Original", "awayMinutes": 12}, headers=H())
    assert r.status_code == 200
    last = client.get("/api/events", headers=H()).json()[-1]
    assert last["event"] == "summary_edited" and last["details"]["chars_before"] == 8 and last["details"]["chars_after"] == 29


def test_create_task_and_validation(env):
    client, _ = env
    r = client.post("/api/tasks", json={"title": "  New lab  ", "kind": "document", "complexity": "simple"}, headers=H())
    assert r.status_code == 201 and r.json()["title"] == "New lab" and r.json()["column"] == "todo"
    assert client.post("/api/tasks", json={"title": ""}, headers=H()).status_code == 422
    r = client.post("/api/tasks", json={"title": "x'); DROP TABLE c3_task;--"}, headers=H())
    assert r.status_code == 201 and len(client.get("/api/tasks", headers=H()).json()) == 6


# ----------------------------------------------------------------- files
def _docx():
    from docx import Document

    d = Document()
    d.add_heading("Database assignment report", 1)
    d.add_paragraph("[x] ER diagram")
    d.add_paragraph("[ ] SQL queries")
    d.add_paragraph("TODO: references")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_file_analysis_reads_docx_and_text_and_refuses_other_types(env):
    client, Session = env
    secret = "SECRET_PASSWORD_123"
    files = [("files", ("report.docx", _docx(), "application/octet-stream")),
             ("files", ("contact.html", f"<h1>Contact</h1>\n<!-- TODO: validate email -->\n{secret}\n".encode(), "text/html")),
             ("files", ("photo.png", b"\x89PNG", "image/png"))]
    r = client.post("/api/files/analyze", files=files, headers=KEY)
    assert r.status_code == 200
    by = {f["name"]: f for f in r.json()["files"]}
    assert by["report.docx"]["supported"] and by["report.docx"]["openCount"] == 2
    assert by["contact.html"]["suggest"]["workingPoint"].endswith("validate email")
    assert by["photo.png"]["supported"] is False
    assert secret not in r.text  # the file text is not returned ...
    with Session() as db:  # ... and nothing is stored
        assert db.scalars(select(m.EventLog)).all() == []


# ----------------------------------------------------------------- export, demo, reset
def test_csv_export_contains_the_events(env):
    client, _ = env
    interrupt(client)
    r = client.get("/api/export.csv", headers=H())
    rows = list(csv.reader(io.StringIO(r.text)))
    assert rows[0] == ["time", "event", "task", "details"]
    assert rows[1][1] == "checkpoint_saved" and json.loads(rows[1][3])["level"] in ("Low", "Medium", "High")
    assert "attachment" in r.headers["content-disposition"]


def test_demo_data_and_reset(env):
    client, _ = env
    assert client.post("/api/demo", headers=H()).status_code == 204
    t = tasks(client)
    assert t["Hotel booking website"]["column"] == "interrupted" and t["Hotel booking website"]["checkpoint"]["level"] == "High"
    assert t["Database assignment report"]["checkpoint"]["level"] == "Low"
    assert client.post("/api/reset", headers=H()).status_code == 204
    assert {x["column"] for x in tasks(client).values()} == {"todo", "doing", "done"}
    assert client.get("/api/events", headers=H()).json() == []


def test_reset_can_be_disabled(env, monkeypatch):
    client, _ = env
    monkeypatch.setenv("C3_ALLOW_RESET", "0")
    assert client.post("/api/reset", headers=H()).status_code == 403


# ----------------------------------------------------------------- the learned summary model and its safety net
def _row(level="High", seed=3):
    rows = generate_split("test", random.Random(seed))
    return next(r for r in rows if r["level"] == level and (r["fields"].get("problem") or "").endswith("?"))


def _camel(f):
    return {"goal": f.get("goal"), "completed": f.get("completed"), "workingPoint": f.get("working_point"),
            "problem": f.get("problem"), "nextStep": f.get("next_step"), "resources": f.get("resources"),
            "openItems": f["open_items"]}


class Fake(LoraGenerator):
    def __init__(self, answer):
        super().__init__(model=object())
        self.answer = answer

    def generate(self, input_text):
        return self.answer


def test_serving_input_has_exactly_the_training_format():
    for r in generate_split("val", random.Random(5))[:20]:
        title = r["input"].split("\n")[0].replace("Task: ", "")
        text, fields = build_model_input(title, r["kind"], _camel(r["fields"]), r["level"], r["minutes"])
        assert text == r["input"]
        assert fields == {**to_training_fields(_camel(r["fields"]))}


def test_a_good_model_answer_is_used():
    r = _row()
    title = r["input"].split("\n")[0].replace("Task: ", "")
    out = generate_summary("lora", title, r["kind"], _camel(r["fields"]), r["level"], r["minutes"], lora=Fake(r["output"]))
    assert out["mode"] == "lora" and out["fallback_reason"] is None and out["text"] == r["output"]


@pytest.mark.parametrize("bad,why", [
    ("Welcome back. You should try a different framework and make sure to add tests.", "advice"),
    ("", "empty"),
    ("You were working on it. Quantum blockchain synergy leverages paradigms.", "words that are not in the checkpoint"),
])
def test_a_bad_model_answer_falls_back_to_the_template(bad, why):
    r = _row()
    title = r["input"].split("\n")[0].replace("Task: ", "")
    out = generate_summary("lora", title, r["kind"], _camel(r["fields"]), r["level"], r["minutes"], lora=Fake(bad))
    assert out["mode"] == "template" and out["fallback_reason"] and "Welcome back" in out["text"]


def test_a_model_that_crashes_does_not_break_the_screen():
    class Boom(Fake):
        def generate(self, input_text):
            raise RuntimeError("out of memory")

    r = _row()
    out = generate_summary("lora", "T", r["kind"], _camel(r["fields"]), r["level"], r["minutes"], lora=Boom(""))
    assert out["mode"] == "template" and "RuntimeError" in out["fallback_reason"]


def test_lora_mode_without_the_model_installed_uses_the_template(env):
    client, _ = env
    t, _ = interrupt(client)
    s = client.get(f"/api/tasks/{t['id']}/summary", params={"mode": "lora"}, headers=H()).json()
    assert s["mode"] == "template" and s["requestedMode"] == "lora" and "not installed" in s["fallbackReason"]


def test_checker_rejects_answers_that_drop_the_students_question():
    r = _row("Medium")
    ok = check_model_output(r["output"], r["input"], r["fields"], r["level"])
    assert ok == []
    from components.c3_checkpointing.model2.textutils import cap, to_you

    # the model output repeats the question in second person, so that is the text to remove
    no_question = r["output"].replace(cap(to_you(r["fields"]["problem"])), "something")
    assert no_question != r["output"]
    assert check_model_output(no_question, r["input"], r["fields"], r["level"]) != []


def test_first_load_with_parallel_requests_seeds_the_demo_tasks_only_once(env):
    import threading

    from components.c3_checkpointing.app import service

    _, Session = env
    errors = []

    def first_request():
        db = Session()
        try:
            service.seed_if_empty(db, "parallel-user")
        except Exception as exc:  # pragma: no cover - reported below
            errors.append(exc)
        finally:
            db.close()

    threads = [threading.Thread(target=first_request) for _ in range(6)]
    for th in threads:
        th.start()
    for th in threads:
        th.join()
    assert not errors
    db = Session()
    titles = [t.title for t in db.scalars(select(m.Task).where(m.Task.user_id == "parallel-user"))]
    db.close()
    assert sorted(titles) == sorted(title for title, *_ in service.SEED_TASKS)
